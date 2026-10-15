"""A real .xlsx export, alongside the existing CSV export (app/ai/export.py).
Same rows, same columns - this isn't the client's real IBIS import format
(that template has never arrived, see CLAUDE.md's "Still open" list), just a
genuinely-formatted Excel workbook (frozen header, real column widths,
autofilter) instead of bare CSV, so there's something better to open today.

Mislabeled "Phase 6.6" in an earlier pass of this module - the plan's real
6.6 is "Audit trail" (who/when/before/after per change, review progress per
project), a different, unbuilt feature (the `reviews` table + this element's
review-history panel already cover "who/when/before/after" per change; there
is no project-level progress view yet). This is actually a first, partial
step toward **7.1 Excel export**, whose full spec is three sheets (summary
per IBIS code, detail with GlobalIds, warnings) matching the client's
template, plus export history (the `exports` table exists in the schema but
nothing writes to it yet) - this ships only the detail sheet, no summary/
warnings sheets, no history row, and not the client's real template. Swap
for the real spec once the template exists - the row-building logic (same
`rows` tuples the CSV builder already takes) doesn't change, only how each
row gets written out and what else gets added around it.
"""

import io

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from app.ai.export import FIELDNAMES, _format_quantities
from app.db.models import Classification, Element, Quantity

COLUMN_WIDTHS = {
    "ifc_file_id": 12,
    "global_id": 24,
    "ifc_class": 18,
    "name": 30,
    "storey": 16,
    "quantities": 40,
    "suggested_ibis_code": 18,
    "suggested_internal_element": 28,
    "confidence": 12,
    "status": 16,
    "reasoning": 50,
}


def _write_header(sheet: Worksheet) -> None:
    for col_index, field in enumerate(FIELDNAMES, start=1):
        cell = sheet.cell(row=1, column=col_index, value=field)
        cell.font = openpyxl.styles.Font(bold=True)
        column_letter = openpyxl.utils.get_column_letter(col_index)
        sheet.column_dimensions[column_letter].width = COLUMN_WIDTHS.get(field, 16)
    sheet.freeze_panes = "A2"


def build_export_xlsx(
    rows: list[tuple[Element, list[Quantity], Classification | None]],
) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Elements"
    _write_header(sheet)

    for row_index, (element, quantities, classification) in enumerate(rows, start=2):
        values = {
            "ifc_file_id": element.ifc_file_id,
            "global_id": element.global_id,
            "ifc_class": element.ifc_class,
            "name": element.name or "",
            "storey": element.storey or "",
            "quantities": _format_quantities(quantities),
            "suggested_ibis_code": (classification.ibis_code if classification else ""),
            "suggested_internal_element": (
                classification.internal_element if classification else ""
            ),
            "confidence": (
                classification.confidence if classification is not None else ""
            ),
            "status": classification.status if classification else "not_classified",
            "reasoning": classification.reasoning if classification else "",
        }
        for col_index, field in enumerate(FIELDNAMES, start=1):
            sheet.cell(row=row_index, column=col_index, value=values[field])

    if rows:
        last_column_letter = openpyxl.utils.get_column_letter(len(FIELDNAMES))
        sheet.auto_filter.ref = f"A1:{last_column_letter}{len(rows) + 1}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
